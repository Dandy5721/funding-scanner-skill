#!/usr/bin/env python3
"""Build an Excel tracker from a list of funding opportunities.

Expected input JSON schema (list of objects):
[
  {
    "title": "string (required)",
    "sponsor": "string (required, e.g. 'NIH/NIMH', 'NSF/CISE', 'CDMRP/PHRP')",
    "mechanism": "string (e.g. 'R01', 'R21', 'BAA', 'Foundation grant')",
    "fit_tier": "string ('high' | 'medium' | 'low' | 'trainee' | 'limited')",
    #            'limited' = strong fit but LIMITED SUBMISSION / institutional nomination;
    #            put the institution's internal deadline in loi_deadline and note it in flags.
    "loi_deadline": "string YYYY-MM-DD or '' if none",
    "full_deadline": "string YYYY-MM-DD (required)",
    "award_amount": "string (e.g. '$2M direct over 5 years', 'Up to $50K')",
    "duration": "string (e.g. '5 years')",
    "eligibility": "string (one-line summary)",
    "summary": "string (1-2 sentence what-it-funds)",
    "fit_notes": "string (why it matches the profile)",
    "flags": "string ('LOI required', 'Limited submission', 'Cost share', etc., comma-separated)",
    "url": "string (direct link to announcement)",
    "status": "string ('open' | 'urgent' | 'closing-soon')"
  },
  ...
]

Usage:
    python build_tracker.py --opportunities opps.json --output tracker.xlsx
    python build_tracker.py --opportunities opps.json --output tracker.xlsx --pi-name "<PI full name>"
"""

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNS = [
    ("Fit", 6),
    ("Title", 50),
    ("Sponsor", 18),
    ("Mechanism", 14),
    ("LOI Deadline", 14),
    ("Full Deadline", 14),
    ("Days Until Deadline", 12),
    ("Award", 24),
    ("Duration", 12),
    ("Eligibility", 30),
    ("Summary", 50),
    ("Fit Notes", 40),
    ("Flags", 24),
    ("URL", 50),
    ("Status", 14),
]

FIT_ORDER = {"high": 0, "medium": 1, "low": 2, "limited": 3, "trainee": 4}
FIT_LABEL = {"high": "★★★", "medium": "★★", "low": "★", "limited": "🏛️", "trainee": "👥"}

URGENT_FILL = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")  # red-100
SOON_FILL = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")    # yellow-100
HEADER_FILL = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")  # blue-800
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", size=14, bold=True)


def parse_date(s):
    """Return a date object from YYYY-MM-DD, or None."""
    if not s:
        return None
    try:
        return datetime.strptime(s.strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def days_until(deadline_str):
    """Days until full_deadline. Negative = past. None if unparseable."""
    d = parse_date(deadline_str)
    if d is None:
        return None
    return (d - date.today()).days


def sort_key(opp):
    """Sort by fit tier, then by deadline ascending."""
    tier = FIT_ORDER.get(opp.get("fit_tier", "low"), 9)
    days = days_until(opp.get("full_deadline", ""))
    if days is None:
        days = 99999
    return (tier, days)


def build_workbook(opportunities, pi_name=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funding Tracker"

    row_offset = 1

    # Title row
    if pi_name:
        ws.cell(row=row_offset, column=1, value=f"Funding Tracker — {pi_name}").font = TITLE_FONT
        ws.merge_cells(
            start_row=row_offset, start_column=1,
            end_row=row_offset, end_column=len(COLUMNS),
        )
        row_offset += 1

    ws.cell(row=row_offset, column=1, value=f"Generated {date.today().isoformat()}").font = BODY_FONT
    ws.merge_cells(
        start_row=row_offset, start_column=1,
        end_row=row_offset, end_column=len(COLUMNS),
    )
    row_offset += 2  # blank row before headers

    header_row = row_offset

    # Column headers
    for idx, (col_name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.row_dimensions[header_row].height = 28
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Sort and write opportunities
    opps_sorted = sorted(opportunities, key=sort_key)

    for i, opp in enumerate(opps_sorted, start=1):
        row = header_row + i
        days = days_until(opp.get("full_deadline", ""))

        # Determine row fill based on urgency
        row_fill = None
        if days is not None:
            if 0 <= days <= 14:
                row_fill = URGENT_FILL
            elif 15 <= days <= 42:
                row_fill = SOON_FILL

        values = [
            FIT_LABEL.get(opp.get("fit_tier", "low"), opp.get("fit_tier", "")),
            opp.get("title", ""),
            opp.get("sponsor", ""),
            opp.get("mechanism", ""),
            opp.get("loi_deadline", ""),
            opp.get("full_deadline", ""),
            days if days is not None else "",
            opp.get("award_amount", ""),
            opp.get("duration", ""),
            opp.get("eligibility", ""),
            opp.get("summary", ""),
            opp.get("fit_notes", ""),
            opp.get("flags", ""),
            opp.get("url", ""),
            opp.get("status", "open"),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_fill is not None:
                cell.fill = row_fill

        # Hyperlink the URL cell
        url_col = next(i for i, (n, _) in enumerate(COLUMNS, start=1) if n == "URL")
        url_cell = ws.cell(row=row, column=url_col)
        if url_cell.value:
            url_cell.hyperlink = url_cell.value
            url_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    # Set autofilter on the data range
    last_col = get_column_letter(len(COLUMNS))
    last_row = header_row + len(opps_sorted)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"

    # Add a Legend sheet
    legend = wb.create_sheet("Legend")
    legend["A1"] = "Funding Tracker Legend"
    legend["A1"].font = TITLE_FONT
    legend.column_dimensions["A"].width = 22
    legend.column_dimensions["B"].width = 70

    legend_rows = [
        ("Fit tier", ""),
        ("★★★", "High fit — keyword + eligibility match, deadline workable"),
        ("★★", "Medium fit — partial keyword match or favored mechanism"),
        ("★", "Worth a glance — broadly eligible but topical fit weak"),
        ("🏛️", "Limited submission — strong fit; needs an internal competition at the institution (LOI col = internal deadline)"),
        ("👥", "Trainee opportunity — F31 / F32 / K99 / T32 / etc."),
        ("", ""),
        ("Row color", ""),
        ("Red", "Deadline within 14 days — urgent"),
        ("Yellow", "Deadline within 15-42 days — start drafting"),
        ("White", "Deadline >42 days out, or no deadline parsed"),
        ("", ""),
        ("Status", ""),
        ("open", "Currently accepting applications"),
        ("urgent", "Deadline imminent — flagged separately"),
        ("closing-soon", "Closes within current cycle"),
    ]
    for i, (a, b) in enumerate(legend_rows, start=3):
        legend.cell(row=i, column=1, value=a).font = (
            Font(name="Arial", size=10, bold=True) if a and not b else BODY_FONT
        )
        legend.cell(row=i, column=2, value=b).font = BODY_FONT
        legend.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--opportunities", required=True, help="Path to opportunities JSON file")
    parser.add_argument("--output", required=True, help="Output xlsx path")
    parser.add_argument("--pi-name", default=None, help="Optional PI name for the title row")
    args = parser.parse_args()

    opps_path = Path(args.opportunities)
    if not opps_path.exists():
        print(f"Error: {opps_path} not found", file=sys.stderr)
        sys.exit(1)

    with opps_path.open() as f:
        opportunities = json.load(f)

    if not isinstance(opportunities, list):
        print("Error: opportunities JSON must be a list of objects", file=sys.stderr)
        sys.exit(1)

    wb = build_workbook(opportunities, pi_name=args.pi_name)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Wrote {output_path} ({len(opportunities)} opportunities)")


if __name__ == "__main__":
    main()
